"""Render a shared poster scene to native PowerPoint shapes and vector PDF."""

from __future__ import annotations

import json
import hashlib
import math
import shutil
from pathlib import Path

import pymupdf
from matplotlib import font_manager
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from poster_data import ROOT, REPO, BLUE, OCHRE, RED, GRAY, DARK, REGIME_NAMES

WIDTH, HEIGHT = 36.0, 24.0
FONT_FAMILY = "DejaVu Sans"
TITLE = "Two-Regime Models for Daily Soil-Moisture\nEstimation in Washington"


def register_fonts() -> None:
    folder = ROOT / "assets/fonts"
    folder.mkdir(parents=True, exist_ok=True)
    for name, weight in [("Poster", "normal"), ("PosterBold", "bold")]:
        source = Path(font_manager.findfont(font_manager.FontProperties(family=FONT_FAMILY, weight=weight)))
        target = folder / source.name
        if not target.exists():
            shutil.copyfile(source, target)
        pdfmetrics.registerFont(TTFont(name, str(target)))


def wrap(text: str, width: float, size: float, bold: bool) -> list[str]:
    font = "PosterBold" if bold else "Poster"
    lines = []
    for paragraph in text.split("\n"):
        line = ""
        for word in paragraph.split():
            candidate = f"{line} {word}".strip()
            if line and pdfmetrics.stringWidth(candidate, font, size) > width * 72:
                lines.append(line)
                line = word
            else:
                line = candidate
        lines.append(line)
    return lines


class Scene:
    def __init__(self):
        self.items = []

    def box(self, x, y, w, h, fill, stroke=None, rounded=True, name=""):
        self.items.append(dict(kind="box", x=x, y=y, w=w, h=h, fill=fill,
                               stroke=stroke, rounded=rounded, name=name))

    def text(self, x, y, w, h, text, size=24, color=DARK, bold=False, align="left", name=""):
        lines = wrap(text, w, size, bold)
        required = len(lines) * size * 1.18 / 72
        if required > h + 0.025:
            raise ValueError(f"Text does not fit ({required:.2f} > {h:.2f} in): {text[:95]}")
        self.items.append(dict(kind="text", x=x, y=y, w=w, h=h, text=text, lines=lines,
                               size=size, color=color, bold=bold, align=align,
                               name=name, required_height=required))

    def image(self, x, y, w, h, path, vector=None, name=""):
        self.items.append(dict(kind="image", x=x, y=y, w=w, h=h,
                               path=Path(path).relative_to(ROOT).as_posix(),
                               vector=Path(vector).relative_to(ROOT).as_posix() if vector else None,
                               name=name))

    def line(self, x1, y1, x2, y2, color=GRAY, width=2, arrow=False):
        self.items.append(dict(kind="line", x1=x1, y1=y1, x2=x2, y2=y2,
                               color=color, width=width, arrow=arrow))

    def section(self, x, y, title):
        self.box(x, y, 11.4, 0.55, RED, "#17672A", name=title)
        self.text(x + 0.10, y + 0.023, 11.2, 0.51, title,
                  size=30, color="#FFFFFF", bold=True, align="center")

    def figure(self, x, y, w, h, name):
        self.image(x, y, w, h, ROOT / "figures" / f"{name}.png",
                   ROOT / "figures" / f"{name}.pdf", name)


class Column:
    """Measured vertical flow; both exports receive exactly the same text positions."""

    def __init__(self, scene, x):
        self.scene, self.x, self.y = scene, x, 3.47

    def section(self, title):
        self.scene.section(self.x, self.y, title)
        self.y += 0.72

    def paragraph(self, text, size=22, bold=False, color=DARK, gap=0.12, name=""):
        height = len(wrap(text, 11.0, size, bold)) * size * 1.18 / 72
        self.scene.text(self.x + 0.2, self.y, 11.0, height, text,
                        size=size, bold=bold, color=color, name=name)
        self.y += height + gap

    def figure(self, name, height, caption, width=10.9):
        self.scene.figure(self.x + (11.4 - width) / 2, self.y, width, height, name)
        self.y += height + 0.07
        self.paragraph(caption, size=20, color="#536875", gap=0.14,
                       name=f"Caption {name}")


def method_diagram(scene, x, y):
    """Compact editable input/router/specialist diagram and hard-assignment equation."""
    inputs = [(0.2, "Satellite", "Moisture, backscatter,\nvegetation, land-surface\ntemperature"),
              (4.0, "Weather", "Rainfall totals and\nrecent rainfall history"),
              (7.8, "Site properties", "Terrain, climate,\nland cover, soil")]
    for offset, title, detail in inputs:
        scene.box(x + offset, y, 3.4, 1.42, "#F0F3F5", stroke="#CED6DC")
        scene.text(x + offset + 0.1, y + 0.09, 3.2, 0.33, title, size=20,
                   bold=True, color="#536875", align="center")
        scene.text(x + offset + 0.08, y + 0.44, 3.24, 0.95, detail, size=19, align="center")
        scene.line(x + offset + 1.7, y + 1.46, x + offset + 1.7, y + 1.63)
    y += 0.30
    scene.line(x + 1.9, y + 1.33, x + 9.5, y + 1.33)
    scene.line(x + 5.7, y + 1.33, x + 5.7, y + 1.48, arrow=True)
    scene.box(x + 0.9, y + 1.52, 9.6, 1.12, "#EEF2F4", stroke="#92A2AD")
    scene.text(x + 1.05, y + 1.61, 9.3, 0.43, "Feature-based grouping", size=26,
               bold=True, align="center")
    scene.text(x + 1.05, y + 2.09, 9.3, 0.37,
               "Standardize → K-means → fixed groups for known stations", size=21, align="center")
    scene.line(x + 5.7, y + 2.69, x + 5.7, y + 2.93)
    scene.line(x + 2.8, y + 2.93, x + 8.6, y + 2.93)
    for i, (offset, color) in enumerate([(0.3, BLUE), (6.1, OCHRE)]):
        scene.line(x + offset + 2.5, y + 2.93, x + offset + 2.5, y + 3.17,
                   color=color, width=3, arrow=True)
        scene.box(x + offset, y + 3.21, 5.0, 1.29, color)
        scene.text(x + offset + 0.12, y + 3.32, 4.76, 0.4,
                   f"Specialist {'AB'[i]}", size=24, bold=True, color="#FFFFFF", align="center")
        scene.text(x + offset + 0.12, y + 3.79, 4.76, 0.36,
                   REGIME_NAMES[i], size=21, color="#FFFFFF", align="center")
        scene.text(x + offset + 0.12, y + 4.18, 4.76, 0.29,
                   "XGBoost regressor", size=17.5, color="#FFFFFF", align="center")
        scene.line(x + offset + 2.5, y + 4.54, x + offset + 2.5, y + 4.76,
                   color=color, width=3)
    scene.line(x + 2.8, y + 4.76, x + 8.6, y + 4.76)
    scene.line(x + 5.7, y + 4.76, x + 5.7, y + 4.95, arrow=True)
    scene.box(x + 0.9, y + 4.99, 9.6, 1.07, DARK)
    scene.text(x + 1.05, y + 5.07, 9.3, 0.41, "Daily soil-moisture estimate", size=25,
               bold=True, color="#FFFFFF", align="center")
    scene.text(x + 1.05, y + 5.58, 9.3, 0.38,
               "ŷ = fᵣ(x),  r = g(s, x) ∈ {A, B}", size=22, color="#FFFFFF", align="center",
               name="Editable hard-routing equation")
    return 6.36


def compose(data: dict) -> Scene:
    register_fonts()
    scene = Scene()
    m = data["metrics"]
    left, middle, right = [Column(scene, x) for x in (0.55, 12.30, 24.05)]
    scene.image(0.72, 0.32, 2.16, 1.9872, ROOT / "assets/su_monogram.png", name="Seattle University monogram")
    scene.image(31.0, 0.76, 4.25, 4.25 * 114 / 269, ROOT / "assets/su_college.png", name="Seattle University college logo")
    scene.text(3.25, 0.38, 27.35, 1.63, TITLE, size=48, bold=True, align="center", name="Poster title")
    scene.text(3.25, 2.02, 27.35, 0.48, "[Author names]  ·  Mentored by [Mentor name, credentials]",
               size=24, align="center", name="Editable author and mentor placeholders")
    scene.text(3.25, 2.52, 27.35, 0.44, "Department of Computer Science, Seattle University",
               size=22, align="center", name="Department affiliation")
    scene.text(3.25, 3.00, 27.35, 0.28, "Summer research showcase  ·  2026",
               size=15, color="#697780", bold=True, align="center")

    left.section("Background and research question")
    left.paragraph("Why soil moisture estimation?", size=27, bold=True, color=RED, gap=0.22)
    left.paragraph(
        "Soil moisture regulates agricultural productivity and drought resilience, making daily "
        "estimates vital for water management. High-resolution modeling tracks moisture recharge "
        "and drydown dynamics, particularly across regions lacking direct measurement infrastructure.")
    left.paragraph(
        "While in-situ ground sensors offer accurate readings, they exist only at limited monitoring "
        "sites. Satellite observations span broader regions but remain intermittent with frequent "
        "data gaps. To generate daily estimates, multi-source modeling combines satellite, weather, "
        "and site data, extending an established pipeline and global XGBoost baseline.")
    question = "Can two specialists improve estimates across Washington's contrasting conditions?"
    question_height = len(wrap(question, 10.56, 24, True)) * 24 * 1.18 / 72
    scene.box(left.x + 0.2, left.y, 11.0, 0.66 + question_height, "#F5F1ED", stroke="#DED5CD")
    scene.text(left.x + 0.42, left.y + 0.12, 10.56, 0.34,
               "Summer research question", size=20, bold=True, color=RED)
    scene.text(left.x + 0.42, left.y + 0.55, 10.56, question_height,
               question, size=24, bold=True)
    left.y += 0.66 + question_height + 0.22
    left.paragraph(
        "Both models are compared under identical benchmarks. For known stations, we assign a "
        "fixed specialist; tests at new sites evaluate dynamic routing under missing satellite inputs.")

    left.section("Study area and data")
    scene.figure(left.x + 0.2, left.y, 7.0, 4.5, "washington_stations")
    scene.text(left.x + 7.55, left.y + 0.25, 3.65, 0.38,
               "Regional contrast", size=22, bold=True, color="#536875")
    map_note = (
        "Stations span coastal areas, the Cascade mountains, and arid eastern Washington, capturing "
        "contrasting hydroclimatic regimes. Leave-one-station-out evaluations test model generalization "
        "across unseen microclimates."
    )
    map_note_height = len(wrap(map_note, 3.65, 20, False)) * 20 * 1.18 / 72
    assert map_note_height <= 3.73, f"Map note too tall: {map_note_height:.2f} in"
    scene.text(left.x + 7.55, left.y + 0.77, 3.65, map_note_height,
               map_note, size=20, color="#536875", name="Study-area interpretation")
    left.y += 4.57
    left.paragraph(
        "Figure 1. Seven Washington training stations and their assigned clusters. Colors designate "
        "cluster regimes, with no geographic boundaries implied.", size=20, color="#536875", gap=0.14,
        name="Caption washington_stations")
    left.paragraph(
        "In-situ ground truth for soil moisture comes from NOAA USCRN and SNOTEL stations. Satellite "
        "inputs comprise SMAP soil moisture, Sentinel-1 radar backscatter from the land surface, "
        "and Sentinel-2 surface reflectance and indices of vegetation greenness and canopy water "
        "content. We use MODIS land-surface temperature, ERA5-Land weather, and terrain, soil, "
        "and climate descriptors.")
    left.paragraph(
        "The data pipeline aligns observations onto a daily grid, imputes missing values, smooths "
        "temporal sequences, and extracts rolling lag features capturing recent weather and satellite signals.")
    for offset, text, fill in [(0.25, "2017 to 2022\nFinal fitting period", "#EDF3F6"),
                               (6.05, "2023 to 2025\nLater-year test", "#F5F1ED")]:
        scene.box(left.x + offset, left.y, 5.1, 1.02, fill)
        scene.text(left.x + offset + 0.15, left.y + 0.13, 4.8, 0.73,
                   text, size=21, bold=True, align="center")
    scene.line(left.x + 5.49, left.y + 0.51, left.x + 5.94, left.y + 0.51, arrow=True)
    left.y += 1.20
    left.paragraph(
        "We train the models with data from 2017 to 2022. We evaluate temporal generalization on "
        "later years (2023 to 2025) and test new Washington sites using leave-one-station-out cross-validation.")

    middle.section("Two-regime modeling")
    middle.paragraph(
        "A single global model enforces one input-to-moisture mapping across all stations. "
        "We divide observations into two environmental regimes and fit a dedicated XGBoost "
        "specialist to each. Each specialist predicts daily volumetric soil moisture at 5 cm "
        "depth (m³/m³).")
    middle.y += method_diagram(scene, middle.x, middle.y) + 0.12
    middle.paragraph(
        "Figure 2. Two-regime routing: each observation is assigned to a specialist fᵣ, where x "
        "denotes input features, s the station, and g the routing function.",
        size=20, color="#536875")
    middle.paragraph(
        "1. Router inputs are imputed using training-set means and standardized to comparable "
        "scales. K-means clustering identifies two centroids, assigning unseen samples to the "
        "nearest centroid using feature inputs alone without target labels.")
    middle.paragraph(
        "2. For known stations, each adopts its majority cluster specialist for temporal consistency. "
        "New stations are routed dynamically each day, falling back to rainfall rules if satellite "
        "features are missing.")
    middle.paragraph(
        "3. Both approaches use identical features, XGBoost settings, evaluation splits, and random "
        "seeds. Each specialist trains solely on its assigned regime. Testing 30 repeated seeds "
        "quantifies training stability, while station-holdout splits evaluate spatial generalization uncertainty.")

    middle.section("Regime interpretation")
    profile = data["profile"]
    wet_month = [float(profile.loc["J_bio_bio13", f"median_cluster{i}"]) for i in range(2)]
    smap = [float(profile.loc["SMAP_sm_pm_interp", f"median_cluster{i}"]) for i in range(2)]
    card_y = middle.y
    for i, (offset, color, detail) in enumerate([
        (0.2, BLUE, "Western and Cascade sites:\nmore rain and higher\nSMAP soil moisture."),
        (6.0, OCHRE, "Spokane and Sourdough Gulch:\nless rain and lower\nSMAP soil moisture."),
    ]):
        x = middle.x + offset
        scene.box(x, card_y, 5.2, 2.83, "#F7F9FA", stroke=color)
        scene.box(x, card_y, 5.2, 0.07, color, rounded=False)
        scene.text(x + 0.18, card_y + 0.18, 4.84, 0.40, REGIME_NAMES[i],
                   size=23, color=color, bold=True, align="center")
        n = int((data["stations"].dominant_cluster == i).sum())
        scene.text(x + 0.18, card_y + 0.65, 4.84, 0.34,
                   f"{n} training stations", size=20, bold=True, align="center")
        scene.text(x + 0.18, card_y + 1.09, 4.84, 1.01, detail, size=19, align="center")
        scene.text(x + 0.18, card_y + 2.08, 4.84, 0.63,
                   f"Wettest-month rain: {wet_month[i]:.0f} mm\nSMAP soil moisture: {smap[i]:.3f} m³/m³",
                   size=18, align="center")
    middle.y += 3.01
    middle.paragraph(
        "The regimes separate primarily along gradients of rainfall and SMAP soil moisture. Because "
        "the wetter group includes lowland and mountain stations, elevation alone does not distinguish "
        "them. Displayed values are input medians; other deployments should select cluster counts "
        "based on local environmental data.")

    right.section("Evaluation and results")
    for offset, key, label in [(0.2, "temporal_difference", "Later-year R² gain"),
                               (6.0, "loso_difference", "Station-holdout R² gain")]:
        scene.text(right.x + offset, right.y, 5.2, 0.56,
                   f"+{m[key]:.4f}", size=32, bold=True, align="center", color=RED)
        scene.text(right.x + offset, right.y + 0.61, 5.2, 0.34,
                   label, size=20, align="center")
    right.y += 1.06
    right.figure("model_comparison", 2.8,
                 "Figure 3. Mean R² across 30 later-year runs and five seeds per held-out station, "
                 "using matched seeds (higher is better).")
    right.paragraph(
        "The two-regime model achieves higher R² across all 30 later-year runs. Spatial station "
        "holdouts also favor specialization on average (+0.0581 R²), though gains vary by regional coverage.")
    right.figure("held_out_station_effects", 3.25,
                 "Figure 4. Held-out station R² differences (four gains and three losses).")
    right.paragraph(
        "Cayuse Pass gains +0.25 R². Conversely, withholding either eastern station removes half the "
        "drier group's data, underscoring sensitivity to sparse regional coverage.")

    right.section("ECE sensor-site evaluation")
    right.paragraph(
        "At five sensor sites deployed by Seattle University's ECE team, missing real-time SMAP "
        "retrievals triggered an automated fallback router using recent rainfall history. Regime "
        "thresholds were calibrated on Washington validation data without target labels.")
    right.figure("ece_routing", 2.1,
                 "Figure 5. Pooled RMSE at ECE sites. The rainfall fallback matches global model error (lower is better).")
    right.paragraph(
        "During the dry summer period (July 20 to August 19, 2026; 150 daily observations), the "
        "fallback assigned all dates to the drier specialist. Full-season multi-regime performance remains untested.")

    right.section("Discussion and next steps")
    right.paragraph(
        "Specialist gains rely on accurate routing and representative training networks. Future work "
        "evaluates multi-season ECE sensor data and strengthens router robustness under intermittent satellite availability.")

    ends = {"left": left.y, "middle": middle.y, "right": right.y}
    print("Draft editable-text word count:", sum(len(item["text"].split())
          for item in scene.items if item["kind"] == "text"))
    print("Measured column ends (inches):", json.dumps(ends))
    assert max(ends.values()) <= 23.55, f"Column extends into print margin: {ends}"
    data["column_ends"] = ends
    data["profile_display"] = {"wettest_month_rain_mm": wet_month, "satellite_moisture_m3m3": smap,
                               "group_station_counts": [int((data["stations"].dominant_cluster == i).sum()) for i in range(2)],
                               "regime_names": REGIME_NAMES}
    print("Displayed regime-profile medians and provenance:")
    print(json.dumps(data["profile_display"], indent=2))
    print(data["manifest"]["inputs"]["profile"]["path"])
    return scene


def hex_rgb(value):
    return tuple(int(value.lstrip("#")[i:i + 2], 16) for i in (0, 2, 4))


def render_pptx(scene: Scene, path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(WIDTH), Inches(HEIGHT)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    for index, item in enumerate(scene.items):
        kind = item["kind"]
        if kind == "box":
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if item["rounded"] else MSO_SHAPE.RECTANGLE,
                                           *[Inches(item[key]) for key in ("x", "y", "w", "h")])
            if item["rounded"]:
                shape.adjustments[0] = 0.08
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(*hex_rgb(item["fill"]))
            if item["stroke"]:
                shape.line.color.rgb = RGBColor(*hex_rgb(item["stroke"]))
                shape.line.width = Pt(1)
            else:
                shape.line.fill.background()
        elif kind == "text":
            shape = slide.shapes.add_textbox(*[Inches(item[key]) for key in ("x", "y", "w", "h")])
            tf = shape.text_frame
            tf.word_wrap = False  # Layout uses measured explicit line breaks, shared with the PDF.
            tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            for line_index, line in enumerate(item["lines"]):
                paragraph = tf.paragraphs[0] if line_index == 0 else tf.add_paragraph()
                paragraph.text = line
                paragraph.space_before = paragraph.space_after = Pt(0)
                paragraph.line_spacing = Pt(item["size"] * 1.18)
                paragraph.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[item["align"]]
                paragraph.font.name = FONT_FAMILY
                paragraph.font.size = Pt(item["size"])
                paragraph.font.bold = item["bold"]
                paragraph.font.color.rgb = RGBColor(*hex_rgb(item["color"]))
        elif kind == "image":
            shape = slide.shapes.add_picture(str(ROOT / item["path"]),
                                             *[Inches(item[key]) for key in ("x", "y", "w", "h")])
        elif kind == "line":
            shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                               *[Inches(item[key]) for key in ("x1", "y1", "x2", "y2")])
            shape.line.color.rgb = RGBColor(*hex_rgb(item["color"]))
            shape.line.width = Pt(item["width"])
            if item["arrow"]:
                end = OxmlElement("a:tailEnd")
                end.set("type", "triangle")
                end.set("w", "med")
                end.set("len", "med")
                shape.line._get_or_add_ln().append(end)
        shape.name = item.get("name") or f"Poster {kind} {index:03}"
    prs.core_properties.title = TITLE.replace("\n", " ")
    prs.core_properties.subject = "Summer research showcase: Washington daily soil-moisture models"
    prs.core_properties.author = "[Author names]"
    prs.core_properties.keywords = "soil moisture, Washington, two regimes, XGBoost, Seattle University"
    prs.save(path)


def render_pdf(scene: Scene, path: Path) -> None:
    base = path.with_suffix(".base.pdf")
    pdf = canvas.Canvas(str(base), pagesize=(WIDTH * 72, HEIGHT * 72), initialFontName="Poster")
    pdf.setTitle(TITLE.replace("\n", " "))
    pdf.setAuthor("[Author names]")
    vector_items = []
    for item in scene.items:
        kind = item["kind"]
        if kind == "box":
            pdf.setFillColorRGB(*(channel / 255 for channel in hex_rgb(item["fill"])))
            if item["stroke"]:
                pdf.setStrokeColorRGB(*(channel / 255 for channel in hex_rgb(item["stroke"])))
            pdf.setLineWidth(1)
            args = (item["x"] * 72, (HEIGHT - item["y"] - item["h"]) * 72, item["w"] * 72, item["h"] * 72)
            if item["rounded"]:
                pdf.roundRect(*args, radius=min(item["w"], item["h"]) * 72 * 0.08,
                              fill=1, stroke=int(bool(item["stroke"])))
            else:
                pdf.rect(*args, fill=1, stroke=int(bool(item["stroke"])))
        elif kind == "text":
            font = "PosterBold" if item["bold"] else "Poster"
            pdf.setFont(font, item["size"])
            pdf.setFillColorRGB(*(channel / 255 for channel in hex_rgb(item["color"])))
            for index, line in enumerate(item["lines"]):
                x = item["x"] * 72
                y = (HEIGHT - item["y"]) * 72 - item["size"] * 0.94 - index * item["size"] * 1.18
                if item["align"] == "center":
                    pdf.drawCentredString(x + item["w"] * 36, y, line)
                elif item["align"] == "right":
                    pdf.drawRightString(x + item["w"] * 72, y, line)
                else:
                    pdf.drawString(x, y, line)
        elif kind == "image":
            if item["vector"]:
                vector_items.append(item)
            else:
                pdf.drawImage(str(ROOT / item["path"]), item["x"] * 72,
                              (HEIGHT - item["y"] - item["h"]) * 72,
                              item["w"] * 72, item["h"] * 72, mask="auto")
        elif kind == "line":
            rgb = tuple(channel / 255 for channel in hex_rgb(item["color"]))
            pdf.setStrokeColorRGB(*rgb)
            pdf.setLineWidth(item["width"])
            x1, x2 = item["x1"] * 72, item["x2"] * 72
            y1, y2 = (HEIGHT - item["y1"]) * 72, (HEIGHT - item["y2"]) * 72
            pdf.line(x1, y1, x2, y2)
            if item["arrow"]:
                angle = math.atan2(y2 - y1, x2 - x1)
                length, spread = 9, 4
                triangle = pdf.beginPath()
                triangle.moveTo(x2, y2)
                for sign in (-1, 1):
                    triangle.lineTo(x2 - length * math.cos(angle) + sign * spread * math.sin(angle),
                                    y2 - length * math.sin(angle) - sign * spread * math.cos(angle))
                triangle.close()
                pdf.setFillColorRGB(*rgb)
                pdf.drawPath(triangle, fill=1, stroke=0)
    pdf.showPage()
    pdf.save()
    with pymupdf.open(base) as document:
        for item in vector_items:
            rect = pymupdf.Rect(item["x"] * 72, item["y"] * 72,
                                (item["x"] + item["w"]) * 72, (item["y"] + item["h"]) * 72)
            with pymupdf.open(ROOT / item["vector"]) as figure:
                document[0].show_pdf_page(rect, figure, keep_proportion=True)
        document.save(path, garbage=4, deflate=True)
    base.unlink()


def validate(scene: Scene, pptx_path: Path, pdf_path: Path) -> dict:
    for item in scene.items:
        if item["kind"] == "line":
            continue
        assert item["x"] >= 0 and item["y"] >= 0
        assert item["x"] + item["w"] <= WIDTH + 0.01
        assert item["y"] + item["h"] <= HEIGHT + 0.01
    prs = Presentation(pptx_path)
    assert len(prs.slides) == 1
    assert prs.slide_width == Inches(WIDTH) and prs.slide_height == Inches(HEIGHT)
    placeholders = [shape for shape in prs.slides[0].shapes if shape.has_text_frame and "[Author names]" in shape.text]
    assert placeholders and "[Mentor name, credentials]" in placeholders[0].text
    assert any(shape.name == "Editable hard-routing equation" and shape.has_text_frame
               for shape in prs.slides[0].shapes)
    with pymupdf.open(pdf_path) as document:
        assert len(document) == 1
        assert tuple(document[0].rect) == (0.0, 0.0, WIDTH * 72, HEIGHT * 72)
        text = document[0].get_text()
        prose = " ".join(text.split())
        for station in ["Beaver Pass", "Cayuse Pass", "Darrington", "Paradise", "Quinault", "Spokane"]:
            assert station in text, station
        assert "Sourdough" in text and "Gulch" in text
        for metric in ["0.7798", "0.8118", "0.5797", "0.6379", "0.1674", "0.0578", "0.0586"]:
            assert metric in text, metric
        for label in REGIME_NAMES.values():
            assert text.count(label) == 3, (label, text.count(label))
        assert "Why soil moisture estimation?" in text
        for obsolete in ["Wet / maritime", "Drier / seasonal", "Why daily estimates?",
                         "Sources:", "shared features", "target depth", "doi:"]:
            assert obsolete not in text, obsolete
        assert text.count("5 cm") == 1
        for term in ["Sentinel-1 radar backscatter", "Sentinel-2 surface reflectance",
                     "canopy water content", "land-surface temperature", "SMAP soil moisture"]:
            assert term in prose, term
        assert "—" not in text and "–" not in text
        for outdated in ["Sentinel radar and optical measurements", "supports a routing lesson"]:
            assert outdated not in prose
        for index in range(1, 6):
            assert f"Figure {index}." in text
        font_embedded = []
        for font in document[0].get_fonts(full=True):
            if font[0]:
                extracted = document.extract_font(font[0])
                font_embedded.append(bool(extracted[3]))
        assert font_embedded and all(font_embedded), "PDF contains an unembedded font"
        preview = ROOT / "summer_showcase_preview.png"
        document[0].get_pixmap(matrix=pymupdf.Matrix(0.92, 0.92), alpha=False).save(preview)
        full = ROOT / "artifacts/print_detail.png"
        document[0].get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False).save(full)
        # Count whitespace-separated words, keeping decimal values and units together.
        word_count = len(text.split())
        print(f"Rendered poster word count (including chart labels): {word_count}")
        assert 850 <= word_count <= 950, f"Poster word budget exceeded: {word_count} words"
    report = {"slides": 1, "pdf_pages": 1, "dimensions_inches": [WIDTH, HEIGHT],
              "pdf_fonts_embedded": all(font_embedded), "author_placeholders_editable": True,
              "scene_items": len(scene.items), "rendered_words": word_count,
              "regime_names": REGIME_NAMES, "equation_editable": True,
              "primary_body_font_pt": 22, "citation_footer_removed": True,
              "charts": "vector PDF, 300 dpi PNG in PowerPoint",
              "preview": preview.name}
    return report


def build(data: dict) -> dict:
    (ROOT / "artifacts").mkdir(parents=True, exist_ok=True)
    scene = compose(data)
    (ROOT / "artifacts/layout.json").write_text(json.dumps(scene.items, indent=2) + "\n")
    (ROOT / "artifacts/profile_display.json").write_text(json.dumps(data["profile_display"], indent=2) + "\n")
    pptx_path = ROOT / "summer_showcase.pptx"
    pdf_path = ROOT / "summer_showcase.pdf"
    render_pptx(scene, pptx_path)
    render_pdf(scene, pdf_path)
    report = validate(scene, pptx_path, pdf_path)
    for relative_path, expected_hash in data["manifest"]["preserved_previous_version"].items():
        assert hashlib.sha256((REPO / relative_path).read_bytes()).hexdigest() == expected_hash, relative_path
    report["previous_version_unchanged"] = True
    report["column_ends_inches"] = data["column_ends"]
    (ROOT / "artifacts/validation.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))
    print("Built summer_showcase.pptx, summer_showcase.pdf, summer_showcase_preview.png")
    return report
