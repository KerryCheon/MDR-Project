"""Copy only the two header logos from the user-supplied poster template."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("template", type=Path)
    args = parser.parse_args()
    presentation = Presentation(args.template)
    images = []
    for shape in presentation.slides[0].shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE or shape.top / 914400 >= 3:
            continue
        try:
            shape.image
        except ValueError:
            continue  # Some non-header reference figures are linked pictures.
        images.append(shape)
    print("Header image positions:")
    for shape in images:
        print(shape.name, tuple(round(value / 914400, 3) for value in
                                (shape.left, shape.top, shape.width, shape.height)))
    left = min(images, key=lambda shape: shape.left)
    right = max(images, key=lambda shape: shape.left)
    assets = ROOT / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    records = []
    for label, shape in [("su_monogram", left), ("su_college", right)]:
        path = assets / f"{label}.{shape.image.ext}"
        path.write_bytes(shape.image.blob)
        records.append({
            "asset": path.relative_to(ROOT).as_posix(),
            "template_shape": shape.name,
            "sha256": hashlib.sha256(shape.image.blob).hexdigest(),
            "pixel_size": list(shape.image.size),
        })
        print(f"Copied {path.relative_to(ROOT)} ({shape.image.size})")
    manifest = {
        "source_filename": args.template.name,
        "source_sha256": hashlib.sha256(args.template.read_bytes()).hexdigest(),
        "source_dimensions_inches": [presentation.slide_width / 914400,
                                      presentation.slide_height / 914400],
        "purpose": "Seattle University branding supplied by the user for this poster",
        "assets": records,
    }
    (assets / "branding_provenance.json").write_text(json.dumps(manifest, indent=2) + "\n")


if __name__ == "__main__":
    main()
