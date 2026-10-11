"""Render a shared poster scene to native PowerPoint shapes and vector PDF."""

from __future__ import annotations

import json
import math
import shutil
from pathlib import Path

import pymupdf
from matplotlib import font_manager
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from poster_data import ROOT, BLUE, OCHRE, RED, GRAY, DARK

WIDTH, HEIGHT = 36.0, 24.0
FONT_FAMILY = "DejaVu Sans"
TITLE = "Two-Regime Models for Daily Soil-Moisture\nEstimation in Washington"


def register_fonts() -> None:
    folder = ROOT / "assets/fonts"
    folder.mkdir(parents=True, exist_ok=True)
    for name, weight in [("Poster", "normal"), ("PosterBold", "bold")]:
        source = Path(font_manager.findfont(font_manager.FontProperties(family=FONT_FAMILY, weight=weight)))
        target = folder / source.name
        if not target.exists():
            shutil.copyfile(source, target)
        pdfmetrics.registerFont(TTFont(name, str(target)))


def wrap(text: str, width: float, size: float, bold: bool) -> list[str]:
    font = "PosterBold" if bold else "Poster"
    lines = []
    for paragraph in text.split("\n"):
        line = ""
        for word in paragraph.split():
            candidate = f"{line} {word}".strip()
            if line and pdfmetrics.stringWidth(candidate, font, size) > width * 72:
                lines.append(line)
                line = word
            else:
                line = candidate
        lines.append(line)
    return lines


class Scene:
    def __init__(self):
        self.items = []

    def box(self, x, y, w, h, fill, stroke=None, rounded=True, name=""):
        self.items.append(dict(kind="box", x=x, y=y, w=w, h=h, fill=fill,
                               stroke=stroke, rounded=rounded, name=name))

    def text(self, x, y, w, h, text, size=24, color=DARK, bold=False, align="left", name=""):
        lines = wrap(text, w, size, bold)
        required = len(lines) * size * 1.18 / 72
        if required > h + 0.025:
            raise ValueError(f"Text does not fit ({required:.2f} > {h:.2f} in): {text[:95]}")
        self.items.append(dict(kind="text", x=x, y=y, w=w, h=h, text=text, lines=lines,
                               size=size, color=color, bold=bold, align=align,
                               name=name, required_height=required))

    def image(self, x, y, w, h, path, vector=None, name=""):
        self.items.append(dict(kind="image", x=x, y=y, w=w, h=h,
                               path=Path(path).relative_to(ROOT).as_posix(),
                               vector=Path(vector).relative_to(ROOT).as_posix() if vector else None,
                               name=name))

    def line(self, x1, y1, x2, y2, color=GRAY, width=2, arrow=False):
        self.items.append(dict(kind="line", x1=x1, y1=y1, x2=x2, y2=y2,
                               color=color, width=width, arrow=arrow))

    def section(self, x, y, title):
        self.box(x, y, 11.4, 0.55, RED, "#17672A", name=title)
        self.text(x + 0.10, y + 0.023, 11.2, 0.51, title,
                  size=30, color="#FFFFFF", bold=True, align="center")

    def figure(self, x, y, w, h, name):
        self.image(x, y, w, h, ROOT / "figures" / f"{name}.png",
                   ROOT / "figures" / f"{name}.pdf", name)


def compose(data: dict) -> Scene:
    register_fonts()
    scene = Scene()
    m = data["metrics"]
    left, middle, right = 0.55, 12.30, 24.05
    scene.image(0.72, 0.32, 2.16, 1.9872, ROOT / "assets/su_monogram.png", name="Seattle University monogram")
    scene.image(31.0, 0.76, 4.25, 4.25 * 114 / 269, ROOT / "assets/su_college.png", name="Seattle University college logo")
    scene.text(3.25, 0.38, 27.35, 1.63, TITLE, size=48, bold=True, align="center", name="Poster title")
    scene.text(3.25, 2.02, 27.35, 0.48, "[Author names]  ·  Mentored by [Mentor name, credentials]",
               size=24, align="center", name="Editable author and mentor placeholders")
    scene.text(3.25, 2.52, 27.35, 0.44, "Department of Computer Science, Seattle University",
               size=22, align="center", name="Department affiliation")
    scene.text(3.25, 3.00, 27.35, 0.28, "SUMMER RESEARCH SHOWCASE  ·  2026",
               size=15, color="#697780", bold=True, align="center")

    # Left: the project, then its study stations. The map remains a small contextual figure.
    scene.section(left, 3.47, "Project Overview")
    scene.text(left + 0.2, 4.25, 11.0, 0.88, "Why daily estimates?", size=27, bold=True, color=RED)
    scene.text(left + 0.2, 5.40, 11.0, 1.70,
               "Soil moisture informs irrigation and drought monitoring. Ground sensors provide local measurements; "
               "satellites cover larger areas but leave gaps in daily observations.", size=24)
    scene.text(left + 0.2, 7.34, 11.0, 1.28,
               "We combine satellite, weather, and site data for daily soil-moisture estimation. "
               "Earlier work established the pipeline and a single global predictor.", size=24)
    scene.box(left + 0.2, 8.91, 11.0, 1.76, "#F5F1ED", stroke="#DED5CD")
    scene.text(left + 0.42, 9.08, 10.56, 0.4, "THE SUMMER QUESTION", size=20, bold=True, color=RED)
    scene.text(left + 0.42, 9.59, 10.56, 0.94,
               "Can two specialists improve estimates across Washington's contrasting conditions?", size=25, bold=True)

    scene.section(left, 11.09, "Study Area & Data")
    scene.figure(left + 2.2, 11.95, 7, 4.5, "washington_stations")
    scene.text(left + 0.4, 16.69, 10.6, 0.97,
               "Station colors show fitted memberships across western Washington, the Cascades, and the east.", size=20)
    for offset, value, label in [(0.25, str(m["stations"]), "public stations"),
                                 (4.0, str(m["features"]), "shared features"),
                                 (7.75, "5 cm", "target depth")]:
        scene.box(left + offset, 18.00, 3.35, 1.45, "#F0F3F5")
        scene.text(left + offset + 0.1, 18.12, 3.15, 0.64, value, size=36, bold=True, align="center", color=BLUE)
        scene.text(left + offset + 0.1, 18.92, 3.15, 0.36, label, size=19, align="center")
    scene.text(left + 0.2, 19.86, 11.0, 1.14,
               "Satellite, weather, and site features are aligned to daily observations; public in-situ sensors supply training targets.", size=23)
    scene.box(left + 0.25, 21.38, 5.10, 1.02, "#EDF3F6")
    scene.box(left + 6.05, 21.38, 5.10, 1.02, "#F5F1ED")
    scene.text(left + 0.4, 21.50, 4.80, 0.70, "2017–2022\nFit period", size=21, bold=True, align="center")
    scene.text(left + 6.2, 21.50, 4.80, 0.70, "2023–2025\nLater-year test", size=21, bold=True, align="center")
    scene.line(left + 5.49, 21.89, left + 5.94, 21.89, arrow=True)
    scene.text(left + 0.2, 22.66, 11.0, 0.40,
               "Target: daily soil moisture (m³/m³).", size=19)

    # Middle: native editable shapes explain the routing decision.
    scene.section(middle, 3.47, "Two-Regime Method")
    scene.text(middle + 0.2, 4.26, 11.0, 1.20,
               "Group observations by their inputs, then fit an XGBoost specialist to each group's conditions.", size=24)
    inputs = [(middle + 0.20, "SATELLITE", "Moisture, radar,\nvegetation, heat"),
              (middle + 4.0, "WEATHER", "Rainfall and\nrecent history"),
              (middle + 7.8, "SITE", "Terrain, climate,\nland cover")]
    for x, title, detail in inputs:
        scene.box(x, 5.96, 3.4, 1.52, "#F0F3F5", stroke="#CED6DC")
        scene.text(x + 0.1, 6.12, 3.2, 0.35, title, size=20, bold=True, color="#536875", align="center")
        scene.text(x + 0.08, 6.64, 3.24, 0.69, detail, size=20, align="center")
        scene.line(x + 1.7, 7.54, x + 1.7, 7.98, width=2)
    scene.line(middle + 1.90, 7.98, middle + 9.50, 7.98, width=2)
    scene.line(middle + 5.70, 7.98, middle + 5.70, 8.40, width=2, arrow=True)
    scene.box(middle + 1.40, 8.50, 8.60, 1.55, "#EEF2F4", stroke="#92A2AD")
    scene.text(middle + 1.55, 8.63, 8.30, 0.47, "Feature-based grouping", size=27, bold=True, align="center")
    scene.text(middle + 1.55, 9.23, 8.30, 0.39, "K-means selects a specialist", size=21, align="center")
    scene.line(middle + 5.70, 10.09, middle + 5.70, 10.47, width=2)
    scene.line(middle + 2.8, 10.47, middle + 8.6, 10.47, width=2)
    scene.line(middle + 2.8, 10.47, middle + 2.8, 10.92, color=BLUE, width=3, arrow=True)
    scene.line(middle + 8.6, 10.47, middle + 8.6, 10.92, color=OCHRE, width=3, arrow=True)
    for offset, color, label, detail in [
        (0.30, BLUE, "Specialist A", "Wet / maritime\ninput conditions"),
        (6.10, OCHRE, "Specialist B", "Drier / seasonal\ninput conditions"),
    ]:
        scene.box(middle + offset, 11.06, 5.0, 2.20, color)
        scene.text(middle + offset + 0.15, 11.25, 4.7, 0.47, label, size=28,
                   bold=True, color="#FFFFFF", align="center")
        scene.text(middle + offset + 0.15, 11.94, 4.7, 0.82, detail, size=24,
                   color="#FFFFFF", align="center")
        scene.text(middle + offset + 0.15, 12.87, 4.7, 0.32, "XGBoost regressor", size=18,
                   color="#FFFFFF", align="center")
    scene.line(middle + 2.8, 13.31, middle + 2.8, 13.70, color=BLUE, width=3)
    scene.line(middle + 8.6, 13.31, middle + 8.6, 13.70, color=OCHRE, width=3)
    scene.line(middle + 2.8, 13.70, middle + 8.6, 13.70, width=2)
    scene.line(middle + 5.7, 13.70, middle + 5.7, 14.15, width=2, arrow=True)
    scene.box(middle + 1.2, 14.25, 9.0, 1.13, "#202B35")
    scene.text(middle + 1.4, 14.43, 8.6, 0.49, "Daily soil-moisture estimate", size=27,
               bold=True, color="#FFFFFF", align="center")
    scene.text(middle + 1.4, 15.01, 8.6, 0.31, "One observation → one specialist → one estimate", size=18,
               color="#FFFFFF", align="center")
    scene.text(middle + 0.3, 15.69, 10.8, 0.71,
               "Known stations keep one specialist; new stations are assigned using their features.", size=21)
    scene.text(middle + 0.3, 16.48, 10.8, 0.69,
               "Grouping uses no observed in-situ targets.", size=20, color="#536875")

    scene.section(middle, 17.46, "What the Regimes Capture")
    profile = data["profile"]
    wet_month = [float(profile.loc["J_bio_bio13", f"median_cluster{i}"]) for i in range(2)]
    smap = [float(profile.loc["SMAP_sm_pm_interp", f"median_cluster{i}"]) for i in range(2)]
    for i, (offset, color, title, detail) in enumerate([
        (0.20, BLUE, "Wet / maritime", "Western Washington:\ncoastal and Cascade\nconditions together."),
        (6.00, OCHRE, "Drier / seasonal", "Eastern Washington:\nless rain and lower\nsatellite moisture."),
    ]):
        x = middle + offset
        scene.box(x, 18.30, 5.2, 3.54, "#F7F9FA", stroke=color)
        scene.box(x, 18.30, 5.2, 0.09, color, rounded=False)
        scene.text(x + 0.18, 18.53, 4.84, 0.49, title, size=25, color=color, bold=True, align="center")
        n = int((data["stations"].dominant_cluster == i).sum())
        scene.text(x + 0.18, 19.14, 4.84, 0.35, f"{n} training stations", size=20, bold=True, align="center")
        scene.text(x + 0.18, 19.70, 4.84, 1.03, detail, size=20, align="center")
        scene.text(x + 0.18, 20.94, 4.84, 0.74,
                   f"Wettest-month rain: {wet_month[i]:.0f} mm\nSatellite moisture: {smap[i]:.3f} m³/m³", size=17.5, align="center")
    scene.text(middle + 0.25, 22.11, 10.9, 0.94,
               "Group medians describe this Washington sample. Choose regime counts using local data.", size=19)

    # Right: comparable evidence, station differences, and the sensor-site lesson.
    scene.section(right, 3.47, "Results")
    for offset, key, label in [(0.2, "temporal_difference", "Later years"),
                               (6.0, "loso_difference", "Held-out stations")]:
        scene.text(right + offset, 4.20, 5.2, 0.64, f"+{m[key]:.4f} R²", size=34,
                   bold=True, align="center", color=RED)
        scene.text(right + offset, 4.96, 5.2, 0.37, label, size=21, align="center")
    scene.figure(right + 0.25, 5.67, 10.9, 3.0, "model_comparison")
    scene.text(right + 0.25, 8.98, 10.9, 0.95,
               "Matched inputs and settings: 30 later-year runs; five per held-out station. "
               "Two-regime R² leads in every later-year run.", size=19.5)
    scene.text(right + 0.25, 10.10, 10.9, 0.45, "Where does specialization help?", size=25, bold=True)
    scene.figure(right + 0.25, 10.78, 10.9, 3.55, "held_out_station_effects")
    scene.text(right + 0.25, 14.46, 10.9, 0.72,
               "Gains at four stations; losses at three. Eastern holdouts reduce representation of the smaller regime.", size=19)

    scene.section(right, 15.49, "New Sensor Sites: Routing Matters")
    scene.text(right + 0.25, 16.29, 10.9, 0.91,
               "Missing SMAP inputs triggered automatic routing based on available precipitation history at five new ECE sites.", size=21)
    scene.figure(right + 0.25, 17.50, 10.9, 2.15, "ece_routing")
    scene.text(right + 0.25, 19.85, 10.9, 0.99,
               "150 observations, July–August 2026; one fallback specialist. Full-season evaluation remains.", size=19)
    scene.section(right, 21.01, "Takeaways")
    scene.text(right + 0.25, 21.79, 10.9, 1.13,
               "Specialization helps when routing is reliable and training sites represent new conditions. "
               "Next: evaluate more stations and seasons with the ECE team's sensors.", size=22)

    scene.line(0.55, 23.22, 35.45, 23.22, color="#CBD3D8", width=1)
    scene.text(0.7, 23.31, 35.0, 0.48,
               "Sources: [1] Balkovec et al., IEEE AIIoT (2026), doi:10.1109/AIIoT68874.2026.11569136.  "
               "[2] MDR Project, technical report 1.5 (2026).\n"
               "In-situ data: NOAA USCRN and USDA ISMN SNOTEL. New sensor-site evaluation in collaboration with the Seattle University ECE team.",
               size=12, color="#536875", name="Sources and acknowledgment")
    # Add all diagram/profile statistics to the audited numerical record.
    data["profile_display"] = {"wettest_month_rain_mm": wet_month, "satellite_moisture_m3m3": smap,
                               "group_station_counts": [int((data["stations"].dominant_cluster == i).sum()) for i in range(2)]}
    print("Displayed regime-profile medians and provenance:")
    print(json.dumps(data["profile_display"], indent=2))
    print(data["manifest"]["inputs"]["profile"]["path"])
    return scene


def hex_rgb(value):
    return tuple(int(value.lstrip("#")[i:i + 2], 16) for i in (0, 2, 4))


def render_pptx(scene: Scene, path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(WIDTH), Inches(HEIGHT)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    for index, item in enumerate(scene.items):
        kind = item["kind"]
        if kind == "box":
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if item["rounded"] else MSO_SHAPE.RECTANGLE,
                                           *[Inches(item[key]) for key in ("x", "y", "w", "h")])
            if item["rounded"]:
                shape.adjustments[0] = 0.08
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(*hex_rgb(item["fill"]))
            if item["stroke"]:
                shape.line.color.rgb = RGBColor(*hex_rgb(item["stroke"]))
                shape.line.width = Pt(1)
            else:
                shape.line.fill.background()
        elif kind == "text":
            shape = slide.shapes.add_textbox(*[Inches(item[key]) for key in ("x", "y", "w", "h")])
            tf = shape.text_frame
            tf.word_wrap = False  # Layout uses measured explicit line breaks, shared with the PDF.
            tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            for line_index, line in enumerate(item["lines"]):
                paragraph = tf.paragraphs[0] if line_index == 0 else tf.add_paragraph()
                paragraph.text = line
                paragraph.space_before = paragraph.space_after = Pt(0)
                paragraph.line_spacing = Pt(item["size"] * 1.18)
                paragraph.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[item["align"]]
                paragraph.font.name = FONT_FAMILY
                paragraph.font.size = Pt(item["size"])
                paragraph.font.bold = item["bold"]
                paragraph.font.color.rgb = RGBColor(*hex_rgb(item["color"]))
        elif kind == "image":
            shape = slide.shapes.add_picture(str(ROOT / item["path"]),
                                             *[Inches(item[key]) for key in ("x", "y", "w", "h")])
        elif kind == "line":
            shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                               *[Inches(item[key]) for key in ("x1", "y1", "x2", "y2")])
            shape.line.color.rgb = RGBColor(*hex_rgb(item["color"]))
            shape.line.width = Pt(item["width"])
            if item["arrow"]:
                end = OxmlElement("a:tailEnd")
                end.set("type", "triangle")
                end.set("w", "med")
                end.set("len", "med")
                shape.line._get_or_add_ln().append(end)
        shape.name = item.get("name") or f"Poster {kind} {index:03}"
    prs.core_properties.title = TITLE.replace("\n", " ")
    prs.core_properties.subject = "Summer research showcase: Washington daily soil-moisture models"
    prs.core_properties.author = "[Author names]"
    prs.core_properties.keywords = "soil moisture, Washington, two regimes, XGBoost, Seattle University"
    prs.save(path)


def render_pdf(scene: Scene, path: Path) -> None:
    base = path.with_suffix(".base.pdf")
    pdf = canvas.Canvas(str(base), pagesize=(WIDTH * 72, HEIGHT * 72), initialFontName="Poster")
    pdf.setTitle(TITLE.replace("\n", " "))
    pdf.setAuthor("[Author names]")
    vector_items = []
    for item in scene.items:
        kind = item["kind"]
        if kind == "box":
            pdf.setFillColorRGB(*(channel / 255 for channel in hex_rgb(item["fill"])))
            if item["stroke"]:
                pdf.setStrokeColorRGB(*(channel / 255 for channel in hex_rgb(item["stroke"])))
            pdf.setLineWidth(1)
            args = (item["x"] * 72, (HEIGHT - item["y"] - item["h"]) * 72, item["w"] * 72, item["h"] * 72)
            if item["rounded"]:
                pdf.roundRect(*args, radius=min(item["w"], item["h"]) * 72 * 0.08,
                              fill=1, stroke=int(bool(item["stroke"])))
            else:
                pdf.rect(*args, fill=1, stroke=int(bool(item["stroke"])))
        elif kind == "text":
            font = "PosterBold" if item["bold"] else "Poster"
            pdf.setFont(font, item["size"])
            pdf.setFillColorRGB(*(channel / 255 for channel in hex_rgb(item["color"])))
            for index, line in enumerate(item["lines"]):
                x = item["x"] * 72
                y = (HEIGHT - item["y"]) * 72 - item["size"] * 0.94 - index * item["size"] * 1.18
                if item["align"] == "center":
                    pdf.drawCentredString(x + item["w"] * 36, y, line)
                elif item["align"] == "right":
                    pdf.drawRightString(x + item["w"] * 72, y, line)
                else:
                    pdf.drawString(x, y, line)
        elif kind == "image":
            if item["vector"]:
                vector_items.append(item)
            else:
                pdf.drawImage(str(ROOT / item["path"]), item["x"] * 72,
                              (HEIGHT - item["y"] - item["h"]) * 72,
                              item["w"] * 72, item["h"] * 72, mask="auto")
        elif kind == "line":
            rgb = tuple(channel / 255 for channel in hex_rgb(item["color"]))
            pdf.setStrokeColorRGB(*rgb)
            pdf.setLineWidth(item["width"])
            x1, x2 = item["x1"] * 72, item["x2"] * 72
            y1, y2 = (HEIGHT - item["y1"]) * 72, (HEIGHT - item["y2"]) * 72
            pdf.line(x1, y1, x2, y2)
            if item["arrow"]:
                angle = math.atan2(y2 - y1, x2 - x1)
                length, spread = 9, 4
                triangle = pdf.beginPath()
                triangle.moveTo(x2, y2)
                for sign in (-1, 1):
                    triangle.lineTo(x2 - length * math.cos(angle) + sign * spread * math.sin(angle),
                                    y2 - length * math.sin(angle) - sign * spread * math.cos(angle))
                triangle.close()
                pdf.setFillColorRGB(*rgb)
                pdf.drawPath(triangle, fill=1, stroke=0)
    pdf.showPage()
    pdf.save()
    with pymupdf.open(base) as document:
        for item in vector_items:
            rect = pymupdf.Rect(item["x"] * 72, item["y"] * 72,
                                (item["x"] + item["w"]) * 72, (item["y"] + item["h"]) * 72)
            with pymupdf.open(ROOT / item["vector"]) as figure:
                document[0].show_pdf_page(rect, figure, keep_proportion=True)
        document.save(path, garbage=4, deflate=True)
    base.unlink()


def validate(scene: Scene, pptx_path: Path, pdf_path: Path) -> dict:
    for item in scene.items:
        if item["kind"] == "line":
            continue
        assert item["x"] >= 0 and item["y"] >= 0
        assert item["x"] + item["w"] <= WIDTH + 0.01
        assert item["y"] + item["h"] <= HEIGHT + 0.01
    prs = Presentation(pptx_path)
    assert len(prs.slides) == 1
    assert prs.slide_width == Inches(WIDTH) and prs.slide_height == Inches(HEIGHT)
    placeholders = [shape for shape in prs.slides[0].shapes if shape.has_text_frame and "[Author names]" in shape.text]
    assert placeholders and "[Mentor name, credentials]" in placeholders[0].text
    with pymupdf.open(pdf_path) as document:
        assert len(document) == 1
        assert tuple(document[0].rect) == (0.0, 0.0, WIDTH * 72, HEIGHT * 72)
        text = document[0].get_text()
        for station in ["Beaver Pass", "Cayuse Pass", "Darrington", "Paradise", "Quinault", "Spokane"]:
            assert station in text, station
        assert "Sourdough" in text and "Gulch" in text
        for metric in ["0.7798", "0.8118", "0.5797", "0.6379", "0.1674", "0.0578", "0.0586"]:
            assert metric in text, metric
        font_embedded = []
        for font in document[0].get_fonts(full=True):
            if font[0]:
                extracted = document.extract_font(font[0])
                font_embedded.append(bool(extracted[3]))
        assert font_embedded and all(font_embedded), "PDF contains an unembedded font"
        preview = ROOT / "summer_showcase_preview.png"
        document[0].get_pixmap(matrix=pymupdf.Matrix(0.92, 0.92), alpha=False).save(preview)
        full = ROOT / "artifacts/print_detail.png"
        document[0].get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False).save(full)
        # Count whitespace-separated words, keeping decimal values and units together.
        word_count = len(text.split())
        print(f"Rendered poster word count (including chart labels): {word_count}")
        assert 400 <= word_count <= 500, f"Poster word budget exceeded: {word_count} words"
    report = {"slides": 1, "pdf_pages": 1, "dimensions_inches": [WIDTH, HEIGHT],
              "pdf_fonts_embedded": all(font_embedded), "author_placeholders_editable": True,
              "scene_items": len(scene.items), "rendered_words": word_count,
              "charts": "vector PDF, 300 dpi PNG in PowerPoint",
              "preview": preview.name}
    return report


def build(data: dict) -> dict:
    (ROOT / "artifacts").mkdir(parents=True, exist_ok=True)
    scene = compose(data)
    (ROOT / "artifacts/layout.json").write_text(json.dumps(scene.items, indent=2) + "\n")
    (ROOT / "artifacts/profile_display.json").write_text(json.dumps(data["profile_display"], indent=2) + "\n")
    pptx_path = ROOT / "summer_showcase.pptx"
    pdf_path = ROOT / "summer_showcase.pdf"
    render_pptx(scene, pptx_path)
    render_pdf(scene, pdf_path)
    report = validate(scene, pptx_path, pdf_path)
    (ROOT / "artifacts/validation.json").write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))
    print("Built summer_showcase.pptx, summer_showcase.pdf, summer_showcase_preview.png")
    return report
